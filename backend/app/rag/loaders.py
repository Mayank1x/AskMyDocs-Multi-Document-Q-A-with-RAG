import os
import io
import time
from typing import List
from tenacity import retry, stop_after_attempt, wait_exponential, retry_if_exception_type
from google import genai
from google.genai.errors import APIError
from langchain_core.documents import Document
from app.config import config
import pypdf
import docx
import pptx
import pandas as pd
from app.models import update_document_status

from google import genai
from google.genai import types
import os

def perform_ocr(image_bytes: bytes, ext: str) -> str:
    try:
        api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
        if not api_key:
            return "[OCR Failed: No Gemini API Key provided]"
            
        client = genai.Client(api_key=api_key)
        mime_type = "image/jpeg" if ext.lower() in ["jpg", "jpeg"] else "image/png"
        
        image_part = types.Part.from_bytes(
            data=image_bytes,
            mime_type=mime_type
        )
        
        response = client.models.generate_content(
            model='gemini-3.5-flash',
            contents=[
                image_part,
                "Please extract all readable text from this image exactly as it appears. If the image contains a chart, graph, or diagram, explain it in detail. If there is no text, provide a highly detailed description of the image."
            ]
        )
        return response.text.strip()
    except Exception as e:
        print(f"Gemini OCR failed: {e}")
        return f"[OCR Failed: {str(e)}]"

def perform_image_caption(image_bytes: bytes, ext: str) -> str:
    # Gemini OCR now handles both extraction and captioning in one pass
    return ""

# ----------------- Loaders ----------------- #

def load_pdf(filepath: str, doc_id: int, user_id: str) -> List[Document]:
    documents = []
    reader = pypdf.PdfReader(filepath)
    filename = os.path.basename(filepath)
    total_pages = len(reader.pages)
    
    for i, page in enumerate(reader.pages):
        page_num = i + 1
        update_document_status(doc_id, f"page {page_num} of {total_pages}")
        text = page.extract_text() or ""
        
        # 1. Text extraction
        if len(text.strip()) >= config.OCR_TEXT_THRESHOLD:
            documents.append(Document(
                page_content=text,
                metadata={"doc_id": doc_id, "filename": filename, "page": page_num, "user_id": user_id, "content_type": "text"}
            ))
            
        # 2. Image extraction (for both scanned pages and embedded charts)
        is_scanned_page = len(text.strip()) < config.OCR_TEXT_THRESHOLD
        
        # To prevent massive slowdowns on large PDFs, only OCR images if it's a scanned page, 
        # or if the user explicitly turns on deep image extraction.
        if is_scanned_page or getattr(config, 'DEEP_IMAGE_OCR', False):
            for image_file_object in page.images:
                try:
                    # Avoid tiny icons
                    if len(image_file_object.data) < 15000: # Increased heuristic to 15KB to avoid logos
                        continue
                except:
                    pass
                    
                ext = image_file_object.name.split('.')[-1].lower()
                if ext not in ['png', 'jpg', 'jpeg']:
                    ext = 'png'
                
                ocr_text = perform_ocr(image_file_object.data, ext)
                if len(ocr_text.strip()) > 20 or is_scanned_page:
                    documents.append(Document(
                        page_content=ocr_text,
                        metadata={"doc_id": doc_id, "filename": filename, "page": page_num, "user_id": user_id, "content_type": "ocr"}
                    ))
                    
    return documents

def load_docx(filepath: str, doc_id: int, user_id: str) -> List[Document]:
    documents = []
    filename = os.path.basename(filepath)
    doc = docx.Document(filepath)
    
    # Text
    full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    if full_text:
        documents.append(Document(
            page_content=full_text,
            metadata={"doc_id": doc_id, "filename": filename, "page": 1, "user_id": user_id, "content_type": "text"}
        ))
        
    # Tables
    for t_idx, table in enumerate(doc.tables):
        table_text = ""
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_data:
                table_text += " | ".join(row_data) + "\n"
        if table_text:
            documents.append(Document(
                page_content=table_text,
                metadata={"doc_id": doc_id, "filename": filename, "page": t_idx + 1, "user_id": user_id, "content_type": "table"}
            ))
            
    # Images
    # python-docx stores images in rels
    img_idx = 1
    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            img_data = rel.target_part.blob
            ext = rel.target_ref.split('.')[-1].lower()
            ocr_text = perform_ocr(img_data, ext)
            if ocr_text:
                documents.append(Document(
                    page_content=ocr_text,
                    metadata={"doc_id": doc_id, "filename": filename, "page": img_idx, "user_id": user_id, "content_type": "ocr"}
                ))
            img_idx += 1
            
    return documents

def load_pptx(filepath: str, doc_id: int, user_id: str) -> List[Document]:
    documents = []
    filename = os.path.basename(filepath)
    prs = pptx.Presentation(filepath)
    
    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        slide_text = []
        
        # Text shapes
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                slide_text.append(shape.text)
                
            # Images
            if shape.shape_type == 13: # 13 is msoPicture
                img_data = shape.image.blob
                ext = shape.image.ext
                ocr_text = perform_ocr(img_data, ext)
                if ocr_text:
                    documents.append(Document(
                        page_content=ocr_text,
                        metadata={"doc_id": doc_id, "filename": filename, "page": slide_num, "user_id": user_id, "content_type": "ocr"}
                    ))
        
        if slide_text:
            documents.append(Document(
                page_content="\n".join(slide_text),
                metadata={"doc_id": doc_id, "filename": filename, "page": slide_num, "user_id": user_id, "content_type": "text"}
            ))
            
        # Speaker notes
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                documents.append(Document(
                    page_content=notes,
                    metadata={"doc_id": doc_id, "filename": filename, "page": slide_num, "user_id": user_id, "content_type": "text", "is_notes": True}
                ))
                
    return documents

def load_spreadsheet(filepath: str, doc_id: int, user_id: str) -> List[Document]:
    documents = []
    filename = os.path.basename(filepath)
    
    # Supports both csv and xlsx via pandas
    if filepath.endswith(".csv"):
        df = pd.read_csv(filepath)
    else:
        df = pd.read_excel(filepath)
        
    df = df.fillna("")
    columns = df.columns.tolist()
    
    # Cap rows to prevent massive files blowing up memory
    max_rows = min(len(df), config.MAX_PAGES * 50) 
    
    for i in range(max_rows):
        row_data = df.iloc[i]
        # Format: "Column1: Value1 | Column2: Value2"
        row_text = " | ".join([f"{col}: {row_data[col]}" for col in columns if str(row_data[col]).strip()])
        if row_text:
            documents.append(Document(
                page_content=row_text,
                metadata={"doc_id": doc_id, "filename": filename, "page": i + 1, "user_id": user_id, "content_type": "table"}
            ))
            
    return documents

def load_text(filepath: str, doc_id: int, user_id: str) -> List[Document]:
    filename = os.path.basename(filepath)
    with open(filepath, "r", encoding="utf-8") as f:
        text = f.read()
        
    content_type = "markdown" if filepath.endswith(".md") else "text"
    
    return [Document(
        page_content=text,
        metadata={"doc_id": doc_id, "filename": filename, "page": 1, "user_id": user_id, "content_type": content_type}
    )]

def load_image(filepath: str, doc_id: int, user_id: str) -> List[Document]:
    filename = os.path.basename(filepath)
    ext = filename.split('.')[-1].lower()
    
    with open(filepath, "rb") as f:
        img_data = f.read()
        
    ocr_text = perform_ocr(img_data, ext)
    documents = []
    
    if ocr_text:
        documents.append(Document(
            page_content=ocr_text,
            metadata={"doc_id": doc_id, "filename": filename, "page": 1, "user_id": user_id, "content_type": "ocr"}
        ))
        
    caption = perform_image_caption(img_data, ext)
    if caption:
        documents.append(Document(
            page_content=caption,
            metadata={"doc_id": doc_id, "filename": filename, "page": 1, "user_id": user_id, "content_type": "image_caption"}
        ))
        
    return documents

def route_loader(filepath: str, doc_id: int, user_id: str) -> List[Document]:
    """Route file to the correct loader based on extension."""
    ext = filepath.split('.')[-1].lower()
    
    if ext == "pdf":
        return load_pdf(filepath, doc_id, user_id)
    elif ext == "docx":
        return load_docx(filepath, doc_id, user_id)
    elif ext == "pptx":
        return load_pptx(filepath, doc_id, user_id)
    elif ext in ["csv", "xlsx"]:
        return load_spreadsheet(filepath, doc_id, user_id)
    elif ext in ["txt", "md", "html"]:
        return load_text(filepath, doc_id, user_id)
    elif ext in ["png", "jpg", "jpeg", "webp"]:
        return load_image(filepath, doc_id, user_id)
    else:
        raise ValueError(f"Unsupported file extension: {ext}")
